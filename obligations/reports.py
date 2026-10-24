"""S4 scoped read adapters. No fake verification or notification success."""
from datetime import timedelta
from io import BytesIO

from django.core.exceptions import ValidationError
from django.utils import timezone

from portal.models import AuditEvent
from portal.permissions import scoped_instruments
from .kpis import AcceptedSuccessor, KpiEvent, ProcessSnapshot, RenewalSnapshot, calculate_kpis
from .services import wire_money

REPORTS = {'pending_verifications', 'upcoming_maturities', 'missing_documents', 'intervention', 'released_register'}


def report_rows(actor, kind, *, as_of=None):
    if kind not in REPORTS:
        raise ValidationError({'kind': 'Unsupported report.'})
    when = as_of or timezone.localdate()
    records = scoped_instruments(actor).select_related('institution').prefetch_related('documents', 'observations')
    rows = []
    for h in records.order_by('pk'):
        documents = list(h.documents.all())
        reasons = []
        if kind == 'pending_verifications':
            if h.status != 'SUBMITTED':
                continue
            reasons = ['Awaiting officer review; independent verification state unavailable']
        elif kind == 'upcoming_maturities':
            if not h.is_current or not when <= h.maturity_date <= when + timedelta(days=90):
                continue
            reasons = ['Current holding matures within 90 days']
        elif kind == 'missing_documents':
            if any(d.kind == 'FDR' for d in documents):
                continue
            reasons = ['FDR supporting document missing; other typed requirements unavailable']
        elif kind == 'intervention':
            if h.is_current and h.maturity_date < when:
                reasons.append('Renewal overdue')
            if any(o.severity == 'BLOCKING' and o.resolved_at is None for o in h.observations.all()):
                reasons.append('Unresolved blocking observation')
            if not reasons:
                continue
        elif h.released_at is None:
            continue
        rows.append({'instrument_id': h.pk, 'institution': h.institution.name, 'reference': h.reference,
            'amount_paise': h.amount_paise, 'maturity_date': h.maturity_date.isoformat(),
            'status': h.status, 'released_at': h.released_at.isoformat() if h.released_at else None,
            'reason': '; '.join(reasons)})
    return wire_money({'kind': kind, 'as_of': when.isoformat(), 'simulation_label': 'SIMULATED',
        'limitations': ['No independent bank-verification adapter',
            'Typed release/checklist document adapter unavailable',
            'Obligation/SLA/discrepancy intervention adapters unavailable'], 'rows': rows})


def xlsx_bytes(report):
    from openpyxl import Workbook
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'SIMULATED report'
    fields = ['instrument_id', 'institution', 'reference', 'amount_paise', 'maturity_date', 'status', 'released_at', 'reason']
    sheet.append(['SIMULATED — not independently bank-verified'])
    sheet.append(['Report', report['kind'], 'As of', report['as_of']])
    for limitation in report['limitations']:
        sheet.append(['Limitation', limitation])
    sheet.append(fields)
    for row in report['rows']:
        # Explicit string cell type prevents formula injection, including leading whitespace.
        sheet.append([ILLEGAL_CHARACTERS_RE.sub('\ufffd', row[key]) if isinstance(row[key], str)
                      else row[key] for key in fields])
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = 's'
    sheet.freeze_panes = 'A7'
    sheet.auto_filter.ref = f'A6:H{max(6, sheet.max_row)}'
    for column in ['A', 'D', 'F']:
        sheet.column_dimensions[column].width = 22
    for column in ['B', 'C', 'E', 'G', 'H']:
        sheet.column_dimensions[column].width = 38
    target = BytesIO()
    workbook.save(target)
    return target.getvalue()


def kpi_report(actor, *, window_start, window_end, as_of=None):
    cutoff = as_of or timezone.now()
    holdings = list(scoped_instruments(actor).order_by('pk'))
    ids = [h.pk for h in holdings]
    histories = {pk: [] for pk in ids}
    for event in AuditEvent.objects.filter(instrument_id__in=ids).order_by('created_at', 'pk'):
        histories[event.instrument_id].append(event)
    processes, renewals = [], []
    for h in holdings:
        history = histories[h.pk]
        complete = bool(history and history[0].action in {'DRAFT_CREATED', 'RENEWAL_DRAFT_CREATED'})
        events = tuple(KpiEvent(e.pk, sequence, e.action, e.created_at)
            for sequence, e in enumerate(history, 1) if e.action in {'SUBMITTED', 'CHANGES_REQUESTED', 'ACCEPTED'})
        processes.append(ProcessSnapshot(h.pk, h.institution_id, complete, events))
        # No Y6 verification events are invented: confirmed share remains unavailable.
        accepted = any(e.action == 'ACCEPTED' for e in history)
        if not accepted:
            continue
        successors = [s for s in holdings if s.predecessor_id == h.pk]
        successor = None
        if successors:
            candidate = successors[0]
            accepted_events = [e for e in histories[candidate.pk] if e.action == 'ACCEPTED']
            if accepted_events:
                successor = AcceptedSuccessor(candidate.pk, candidate.institution_id, accepted_events[0].created_at)
        renewals.append(RenewalSnapshot(h.pk, h.institution_id, h.maturity_date,
            None if h.released_at else True, complete, successor))
    return calculate_kpis(processes=tuple(processes), renewals=tuple(renewals),
        scope_institution_ids=tuple(sorted({h.institution_id for h in holdings})),
        window_start=window_start, window_end=window_end, as_of=cutoff, simulation_label='SIMULATED')
